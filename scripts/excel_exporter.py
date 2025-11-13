"""
Excel Exporter - Sistema de Exportação de Dados do Bot
Exporta dados de operações e monitoramento de risco para planilhas Excel.
"""
import sys
from pathlib import Path
sys.path.append(str(Path(__file__).parent.parent))

import pandas as pd
from datetime import datetime
from typing import Dict, List, Optional, Any
import json
import os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from zipfile import BadZipFile
import logging

# Configuração de logging
logging.basicConfig(
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s',
    level=logging.INFO
)
logger = logging.getLogger(__name__)

# Diretórios
OUTPUTS_DIR = 'outputs'
LOGS_DIR = 'logs'


class ExcelExporter:
    """
    Classe responsável por exportar dados de trading para Excel.
    
    Funcionalidades:
    - Exporta operações concluídas para planilha histórica
    - Exporta dados de monitoramento de risco
    - Formatação automática com cores e estilos
    - Atualização incremental (adiciona novos dados sem sobrescrever)
    """
    
    def __init__(self, outputs_dir: str = OUTPUTS_DIR):
        """
        Inicializa o exportador.
        
        Args:
            outputs_dir: Diretório onde os arquivos Excel serão salvos
        """
        self.outputs_dir = outputs_dir
        os.makedirs(outputs_dir, exist_ok=True)
        logger.info(f"✅ Excel Exporter inicializado - Diretório: {outputs_dir}")
    
    def _get_filepath(self, filename: str) -> str:
        """Retorna o caminho completo do arquivo."""
        return os.path.join(self.outputs_dir, filename)
    
    def _apply_header_style(self, ws, max_col: int):
        """
        Aplica estilo ao cabeçalho da planilha.
        
        Args:
            ws: Worksheet do openpyxl
            max_col: Número máximo de colunas
        """
        # Estilo do cabeçalho
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_alignment = Alignment(horizontal="center", vertical="center")
        
        # Aplica estilo à primeira linha (cabeçalho)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
    
    def _apply_cell_formatting(self, ws):
        """
        Aplica formatação condicional às células.
        
        Args:
            ws: Worksheet do openpyxl
        """
        # Cores para lucro/prejuízo
        profit_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        loss_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        
        # Identifica colunas de lucro/prejuízo
        profit_cols = []
        for idx, cell in enumerate(ws[1], 1):
            if cell.value and any(term in str(cell.value).lower() for term in ['lucro', 'retorno', 'pnl']):
                profit_cols.append(idx)
        
        # Aplica cores condicionais
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for col_idx in profit_cols:
                cell = row[col_idx - 1]
                if cell.value is not None:
                    try:
                        value = float(cell.value)
                        if value > 0:
                            cell.fill = profit_fill
                        elif value < 0:
                            cell.fill = loss_fill
                    except (ValueError, TypeError):
                        pass
        
        # Ajusta largura das colunas
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            
            for cell in column:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
    
    def export_trade_closed(self, trade_data: Dict[str, Any], filename: str = "historico_operacoes.xlsx") -> bool:
        """
        Exporta uma operação fechada para o histórico em Excel.
        
        Args:
            trade_data: Dicionário com dados da operação fechada
            filename: Nome do arquivo Excel
            
        Returns:
            True se exportado com sucesso, False caso contrário
        """
        try:
            filepath = self._get_filepath(filename)
            
            # Prepara dados para DataFrame
            df_new = pd.DataFrame([{
                'Trade ID': trade_data.get('trade_id', 'N/A'),
                'Criptomoeda': trade_data.get('cripto', 'N/A'),
                'Data Entrada': trade_data.get('timestamp_entrada', 'N/A'),
                'Preço Entrada': trade_data.get('preco_entrada', 0),
                'Data Saída': trade_data.get('timestamp_saida', datetime.now().strftime('%Y-%m-%d %H:%M:%S')),
                'Preço Saída': trade_data.get('preco_saida', 0),
                'Valor Investido': trade_data.get('valor_investido', 100),
                'Lucro Líquido (USD)': trade_data.get('lucro_liquido', 0),
                'Retorno (%)': trade_data.get('retorno_percentual', 0) * 100,  # Converte para %
                'Duração': trade_data.get('duracao', 'N/A'),
                'Motivo Saída': trade_data.get('saida_por', 'N/A'),
                'Stop Loss Inicial': trade_data.get('preco_stop_loss_inicial', 0),
                'Maior Preço Atingido': trade_data.get('trailing_topo_atual', 0),
                'Status': trade_data.get('status', 'CONCLUIDO')
            }])
            
            # Verifica se arquivo existe
            if os.path.exists(filepath):
                df_existing = pd.DataFrame()
                try:
                    df_existing = pd.read_excel(filepath)
                except (ValueError, BadZipFile) as read_error:
                    backup = f"{filepath}.bak_{datetime.now().strftime('%Y%m%d%H%M%S')}"
                    logger.warning(f"⚠️ Arquivo {filename} corrompido ({read_error}). Fazendo backup em {backup} e recriando.")
                    try:
                        os.replace(filepath, backup)
                    except OSError as backup_error:
                        logger.warning(f"⚠️ Falha ao mover arquivo corrompido: {backup_error}")
                
                if not df_existing.empty:
                    # Verifica se trade_id já existe para evitar duplicatas
                    if 'Trade ID' in df_existing.columns and trade_data.get('trade_id') in df_existing['Trade ID'].values:
                        logger.warning(f"⚠️ Trade {trade_data.get('trade_id')} já existe no histórico")
                        return False
                    df_combined = pd.concat([df_existing, df_new], ignore_index=True)
                else:
                    df_combined = df_new
            else:
                df_combined = df_new
            
            # Ordena por data de entrada (mais recente primeiro)
            df_combined = df_combined.sort_values('Data Entrada', ascending=False)
            
            # Salva no Excel com formatação
            with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
                df_combined.to_excel(writer, sheet_name='Histórico', index=False)
                
                # Aplica formatação
                ws = writer.sheets['Histórico']
                self._apply_header_style(ws, df_combined.shape[1])
                self._apply_cell_formatting(ws)
            
            logger.info(f"✅ Operação {trade_data.get('trade_id')} exportada para {filename}")
            return True
            
        except Exception as e:
            logger.error(f"❌ Erro ao exportar operação: {e}")
            import traceback
            logger.error(traceback.format_exc())
            return False
    
    def export_risk_monitoring(self, monitoring_data: Dict[str, Any], filename: str = "monitoramento_risco.xlsx") -> bool:
        """
        Exporta dados de monitoramento de risco para Excel.
        
        Args:
            monitoring_data: Dicionário com dados do monitoramento
            filename: Nome do arquivo Excel
            
        Returns:
            True se exportado com sucesso, False caso contrário
        """
        try:
            filepath = self._get_filepath(filename)
            
            # Prepara dados para DataFrame
            df_new = pd.DataFrame([{
                'Timestamp': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
                'Símbolo': monitoring_data.get('symbol', 'N/A'),
                'Side': monitoring_data.get('side', 'N/A'),
                'Preço Entrada': monitoring_data.get('entry_price', 0),
                'Preço Atual': monitoring_data.get('current_price', 0),
                'PNL (%)': monitoring_data.get('pnl_percentage', 0) * 100,
                'PNL (USD)': monitoring_data.get('unrealized_pnl', 0),
                'Stop Loss Atual': monitoring_data.get('current_stop_loss', 0),
                'Take Profit': monitoring_data.get('take_profit', 0),
                'Maior Preço': monitoring_data.get('highest_price', 0),
                'Trailing Ativo': 'SIM' if monitoring_data.get('is_trailing_active', False) else 'NÃO',
                'Duração (min)': monitoring_data.get('duration_minutes', 0),
                'Status': monitoring_data.get('status', 'MONITORANDO')
            }])
            
            # Verifica se arquivo existe
            if os.path.exists(filepath):
                df_existing = pd.DataFrame()
                try:
                    df_existing = pd.read_excel(filepath)
                except (ValueError, BadZipFile) as read_error:
                    backup = f"{filepath}.bak_{datetime.now().strftime('%Y%m%d%H%M%S')}"
                    logger.warning(f"⚠️ Arquivo {filename} corrompido ({read_error}). Fazendo backup em {backup} e recriando.")
                    try:
                        os.replace(filepath, backup)
                    except OSError as backup_error:
                        logger.warning(f"⚠️ Falha ao mover arquivo corrompido: {backup_error}")

                if not df_existing.empty:
                    df_combined = pd.concat([df_existing, df_new], ignore_index=True)
                else:
                    df_combined = df_new
            else:
                df_combined = df_new
            
            # Mantém apenas os últimos 1000 registros para não ficar muito grande
            if len(df_combined) > 1000:
                df_combined = df_combined.tail(1000)
            
            # Salva no Excel com formatação
            with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
                df_combined.to_excel(writer, sheet_name='Monitoramento', index=False)
                
                ws = writer.sheets['Monitoramento']
                self._apply_header_style(ws, df_combined.shape[1])
                self._apply_cell_formatting(ws)
            
            logger.info(f"✅ Dados de monitoramento exportados para {filename}")
            return True
            
        except Exception as e:
            logger.error(f"❌ Erro ao exportar monitoramento: {e}")
            import traceback
            logger.error(traceback.format_exc())
            return False
    
    def export_daily_summary(self, filename: str = "resumo_diario.xlsx") -> bool:
        """
        Gera um resumo diário com estatísticas de trading.
        
        Args:
            filename: Nome do arquivo Excel
            
        Returns:
            True se exportado com sucesso, False caso contrário
        """
        try:
            # Carrega dados do histórico
            history_file = self._get_filepath("historico_operacoes.xlsx")
            if not os.path.exists(history_file):
                logger.warning("⚠️ Arquivo de histórico não encontrado")
                return False
            
            try:
                df = pd.read_excel(history_file)
            except (ValueError, BadZipFile) as read_error:
                backup = f"{history_file}.bak_{datetime.now().strftime('%Y%m%d%H%M%S')}"
                logger.warning(f"⚠️ Arquivo de histórico corrompido ({read_error}). Movendo para {backup} e abortando resumo.")
                try:
                    os.replace(history_file, backup)
                except OSError as backup_error:
                    logger.warning(f"⚠️ Falha ao mover histórico corrompido: {backup_error}")
                return False
            
            # Converte datas com formato flexível
            df['Data Entrada'] = pd.to_datetime(df['Data Entrada'], format='mixed', errors='coerce')
            df['Data'] = df['Data Entrada'].dt.date
            
            # Agrupa por data
            daily_summary = df.groupby('Data').agg({
                'Trade ID': 'count',
                'Lucro Líquido (USD)': 'sum',
                'Retorno (%)': 'mean',
                'Valor Investido': 'sum'
            }).reset_index()
            
            daily_summary.columns = ['Data', 'Total Trades', 'Lucro Total (USD)', 'Retorno Médio (%)', 'Capital Investido']
            
            # Adiciona taxa de sucesso
            df['Sucesso'] = df['Lucro Líquido (USD)'] > 0
            success_rate = df.groupby('Data')['Sucesso'].mean() * 100
            daily_summary['Taxa Sucesso (%)'] = success_rate.values
            
            # Ordena por data (mais recente primeiro)
            daily_summary = daily_summary.sort_values('Data', ascending=False)
            
            filepath = self._get_filepath(filename)
            
            # Salva no Excel com formatação
            with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
                daily_summary.to_excel(writer, sheet_name='Resumo Diário', index=False)
                
                ws = writer.sheets['Resumo Diário']
                self._apply_header_style(ws, daily_summary.shape[1])
                self._apply_cell_formatting(ws)
            
            logger.info(f"✅ Resumo diário exportado para {filename}")
            return True
            
        except Exception as e:
            logger.error(f"❌ Erro ao gerar resumo diário: {e}")
            import traceback
            logger.error(traceback.format_exc())
            return False
    
    def export_trades_from_json(self, json_file: str = "logs/trades_principais_em_andamento.json") -> bool:
        """
        Exporta todas as operações concluídas de um arquivo JSON para Excel.
        
        Args:
            json_file: Caminho do arquivo JSON
            
        Returns:
            True se exportado com sucesso, False caso contrário
        """
        try:
            if not os.path.exists(json_file):
                logger.warning(f"⚠️ Arquivo {json_file} não encontrado")
                return False
            
            # Carrega dados do JSON
            with open(json_file, 'r') as f:
                trades_data = json.load(f)
            
            # Filtra apenas trades concluídos
            completed_trades = {
                trade_id: data for trade_id, data in trades_data.items()
                if data.get('status') == 'CONCLUIDO'
            }
            
            if not completed_trades:
                logger.info("ℹ️ Nenhuma operação concluída para exportar")
                return False
            
            logger.info(f"📊 Exportando {len(completed_trades)} operação(ões) concluída(s)")
            
            # Exporta cada trade
            success_count = 0
            for trade_id, trade_data in completed_trades.items():
                if self.export_trade_closed(trade_data):
                    success_count += 1
            
            logger.info(f"✅ {success_count}/{len(completed_trades)} operações exportadas com sucesso")
            return True
            
        except Exception as e:
            logger.error(f"❌ Erro ao exportar trades do JSON: {e}")
            import traceback
            logger.error(traceback.format_exc())
            return False


def export_closed_trade(trade_data: Dict[str, Any]) -> bool:
    """
    Função auxiliar para exportar uma operação fechada.
    Pode ser chamada de qualquer lugar do código.
    
    Args:
        trade_data: Dicionário com dados da operação fechada
        
    Returns:
        True se exportado com sucesso, False caso contrário
    """
    exporter = ExcelExporter()
    return exporter.export_trade_closed(trade_data)


def export_risk_monitor_data(monitoring_data: Dict[str, Any]) -> bool:
    """
    Função auxiliar para exportar dados de monitoramento de risco.
    Pode ser chamada de qualquer lugar do código.
    
    Args:
        monitoring_data: Dicionário com dados do monitoramento
        
    Returns:
        True se exportado com sucesso, False caso contrário
    """
    exporter = ExcelExporter()
    return exporter.export_risk_monitoring(monitoring_data)


def generate_daily_summary() -> bool:
    """
    Função auxiliar para gerar resumo diário.
    
    Returns:
        True se gerado com sucesso, False caso contrário
    """
    exporter = ExcelExporter()
    return exporter.export_daily_summary()


# Exemplo de uso
if __name__ == "__main__":
    # Teste de exportação
    exporter = ExcelExporter()
    
    # Exporta trades do JSON
    print("🔄 Exportando trades do JSON...")
    exporter.export_trades_from_json()
    
    # Gera resumo diário
    print("📊 Gerando resumo diário...")
    exporter.export_daily_summary()
    
    print("✅ Exportação concluída!")
